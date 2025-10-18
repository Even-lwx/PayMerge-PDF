#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
发票合并工具 v5.0 - 简化稳定版
基于 v4.0 拖放版本升级，新增智能数据提取功能
简化版本，专注于稳定性和可靠性
"""

import tkinter as tk
from tkinter import messagebox, filedialog
import os
import sys
import threading
from pathlib import Path
import shutil
import tempfile
import json
import csv
import re
from datetime import datetime
from typing import List, Optional, Tuple, Dict, Any

# 尝试导入拖放功能库
try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    DRAG_DROP_AVAILABLE = True
except ImportError:
    DRAG_DROP_AVAILABLE = False
    # 使用普通的Tk作为备选
    TkinterDnD = tk

# 尝试导入PDF处理库
try:
    import pypdfium2 as pdfium
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False
    pdfium = None

# 尝试导入Excel处理库
try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    openpyxl = None

# 导入原有的合并逻辑
try:
    from merge_invoices_simple import merge_simple
except ImportError:
    def merge_simple(pdf_path, img1_path, img2_path, output_path):
        raise ImportError("找不到合并功能模块")


class SimpleInvoiceMergerV5:
    """v5.0 简化版智能发票合并工具"""
    
    def __init__(self):
        # 使用支持的Tk类型
        if DRAG_DROP_AVAILABLE:
            self.root = TkinterDnD.Tk()
        else:
            self.root = tk.Tk()
            
        self.root.title("发票合并工具 v5.0 - 简化稳定版")
        self.root.geometry("800x700")
        self.root.resizable(True, True)

        # 文件存储
        self.pdf_file = None
        self.image_files = []
        
        # 数据提取结果
        self.extracted_data = None
        
        # CSV文件路径
        self.csv_path = self.get_csv_path()

        # 颜色主题
        self.colors = {
            'bg': '#f5f5f5',
            'drop_zone': '#e8f4fd',
            'drop_zone_active': '#c3e7fd',
            'button': '#4a90e2',
            'success': '#52c41a',
            'error': '#f5222d',
            'warning': '#fa8c16',
            'text': '#333333',
        }

        self.setup_ui()
        self.init_csv_file()

    def get_csv_path(self):
        """安全获取CSV文件路径"""
        try:
            if getattr(sys, 'frozen', False):
                app_dir = os.path.dirname(sys.executable)
            else:
                app_dir = os.path.dirname(os.path.abspath(__file__))
            return os.path.join(app_dir, "发票汇总记录.csv")
        except:
            return "发票汇总记录.csv"

    def init_csv_file(self):
        """初始化CSV汇总文件"""
        try:
            if not os.path.exists(self.csv_path):
                headers = ['发票号码', '开票日期', '金额', '销售方名称', '原文件名', '合并文件名', '处理时间']
                with open(self.csv_path, 'w', encoding='utf-8-sig', newline='') as f:
                    writer = csv.writer(f)
                    writer.writerow(headers)
        except Exception as e:
            print(f"CSV文件初始化失败: {e}")

    def setup_ui(self):
        self.root.configure(bg=self.colors['bg'])

        # 标题
        title_label = tk.Label(
            self.root,
            text="📄 发票合并工具 v5.0 - 简化稳定版",
            font=("微软雅黑", 16, "bold"),
            bg=self.colors['bg'],
            fg=self.colors['text']
        )
        title_label.pack(pady=(20, 10))

        # 功能说明
        info_text = "🎯 拖放或选择文件：1个PDF发票 + 2张图片\n✨ 智能命名：日期_金额_发票号后4位_已合并.pdf"
        info_label = tk.Label(
            self.root,
            text=info_text,
            font=("微软雅黑", 11),
            bg=self.colors['bg'],
            fg='#666666'
        )
        info_label.pack(pady=(0, 15))

        # 主容器
        main_frame = tk.Frame(self.root, bg=self.colors['bg'])
        main_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=(0, 20))

        # 拖放区域或文件选择区域
        self.setup_file_area(main_frame)
        
        # 文件列表
        self.setup_file_list(main_frame)
        
        # 数据显示区域
        self.setup_data_display(main_frame)
        
        # 按钮区域
        self.setup_buttons(main_frame)
        
        # 状态栏
        self.setup_status_bar(main_frame)

    def setup_file_area(self, parent):
        """设置文件选择区域"""
        if DRAG_DROP_AVAILABLE:
            # 拖放区域
            self.drop_zone = tk.Frame(
                parent,
                bg=self.colors['drop_zone'],
                relief=tk.RIDGE,
                borderwidth=3,
                height=120
            )
            self.drop_zone.pack(fill=tk.X, pady=(0, 10))

            self.drop_label = tk.Label(
                self.drop_zone,
                text="🎯 将文件拖放到这里\n支持格式：PDF、JPG、PNG",
                font=("微软雅黑", 14),
                bg=self.colors['drop_zone'],
                fg='#666666'
            )
            self.drop_label.pack(expand=True)
            
            # 设置拖放功能
            try:
                self.drop_zone.drop_target_register(DND_FILES)
                self.drop_zone.dnd_bind('<<DropEnter>>', self.on_drop_enter)
                self.drop_zone.dnd_bind('<<DropLeave>>', self.on_drop_leave)
                self.drop_zone.dnd_bind('<<Drop>>', self.on_drop)
            except:
                pass
        else:
            # 普通文件选择区域
            select_frame = tk.Frame(parent, bg=self.colors['bg'])
            select_frame.pack(fill=tk.X, pady=(0, 10))
            
            tk.Label(
                select_frame,
                text="📁 点击按钮选择文件（需要1个PDF + 2张图片）",
                font=("微软雅黑", 12),
                bg=self.colors['bg'],
                fg=self.colors['text']
            ).pack(pady=10)

    def setup_file_list(self, parent):
        """设置文件列表区域"""
        files_frame = tk.LabelFrame(
            parent,
            text="📁 已选择的文件",
            font=("微软雅黑", 11, "bold"),
            bg=self.colors['bg'],
            fg=self.colors['text']
        )
        files_frame.pack(fill=tk.X, pady=(0, 10))

        self.file_listbox = tk.Listbox(
            files_frame,
            height=3,
            font=("微软雅黑", 10),
            bg='white'
        )
        self.file_listbox.pack(fill=tk.X, padx=10, pady=5)

    def setup_data_display(self, parent):
        """设置数据显示区域"""
        data_frame = tk.LabelFrame(
            parent,
            text="🔍 提取的发票数据和智能命名预览",
            font=("微软雅黑", 11, "bold"),
            bg=self.colors['bg'],
            fg=self.colors['text']
        )
        data_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 10))

        self.data_text = tk.Text(
            data_frame,
            height=12,
            font=("Consolas", 9),
            bg='#f6ffed',
            fg=self.colors['text'],
            state=tk.DISABLED,
            wrap=tk.WORD
        )
        self.data_text.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        # 滚动条
        scrollbar = tk.Scrollbar(self.data_text)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.data_text.config(yscrollcommand=scrollbar.set)
        scrollbar.config(command=self.data_text.yview)

    def setup_buttons(self, parent):
        """设置按钮区域"""
        button_frame = tk.Frame(parent, bg=self.colors['bg'])
        button_frame.pack(fill=tk.X, pady=(0, 10))

        # 左侧按钮
        self.select_btn = tk.Button(
            button_frame,
            text="📁 选择文件",
            font=("微软雅黑", 10),
            bg=self.colors['success'],
            fg='white',
            padx=15,
            pady=8,
            command=self.select_files
        )
        self.select_btn.pack(side=tk.LEFT, padx=(0, 10))

        self.clear_btn = tk.Button(
            button_frame,
            text="🗑️ 清除",
            font=("微软雅黑", 10),
            bg=self.colors['error'],
            fg='white',
            padx=15,
            pady=8,
            command=self.clear_files
        )
        self.clear_btn.pack(side=tk.LEFT, padx=(0, 10))

        # 批量提取按钮
        self.batch_extract_btn = tk.Button(
            button_frame,
            text="📊 批量提取",
            font=("微软雅黑", 10, "bold"),
            bg='#722ed1',
            fg='white',
            padx=15,
            pady=8,
            command=self.batch_extract_invoices
        )
        self.batch_extract_btn.pack(side=tk.LEFT, padx=(0, 10))

        # 右侧按钮
        self.extract_btn = tk.Button(
            button_frame,
            text="🔍 提取数据",
            font=("微软雅黑", 10),
            bg=self.colors['warning'],
            fg='white',
            padx=15,
            pady=8,
            command=self.extract_data,
            state=tk.DISABLED
        )
        self.extract_btn.pack(side=tk.RIGHT, padx=(10, 0))

        self.merge_btn = tk.Button(
            button_frame,
            text="🚀 智能合并",
            font=("微软雅黑", 11, "bold"),
            bg=self.colors['button'],
            fg='white',
            padx=20,
            pady=8,
            command=self.merge_files,
            state=tk.DISABLED
        )
        self.merge_btn.pack(side=tk.RIGHT, padx=(10, 0))

    def setup_status_bar(self, parent):
        """设置状态栏"""
        status_frame = tk.Frame(parent, bg=self.colors['bg'])
        status_frame.pack(fill=tk.X)
        
        self.status_label = tk.Label(
            status_frame,
            text="请选择或拖放文件开始操作",
            font=("微软雅黑", 10),
            bg=self.colors['bg'],
            fg='#666666'
        )
        self.status_label.pack(side=tk.LEFT)
        
        # CSV文件信息
        csv_name = os.path.basename(self.csv_path) if self.csv_path else "发票汇总记录.csv"
        csv_info = tk.Label(
            status_frame,
            text=f"📊 汇总文件: {csv_name}",
            font=("微软雅黑", 9),
            bg=self.colors['bg'],
            fg='#999999'
        )
        csv_info.pack(side=tk.RIGHT)

    # 拖放事件处理
    def on_drop_enter(self, event):
        if DRAG_DROP_AVAILABLE:
            self.drop_zone.config(bg=self.colors['drop_zone_active'])

    def on_drop_leave(self, event):
        if DRAG_DROP_AVAILABLE:
            self.drop_zone.config(bg=self.colors['drop_zone'])

    def on_drop(self, event):
        if DRAG_DROP_AVAILABLE:
            self.drop_zone.config(bg=self.colors['drop_zone'])
            files = self.root.tk.splitlist(event.data)
            self.add_files(files)

    def select_files(self):
        """通过对话框选择文件"""
        files = filedialog.askopenfilenames(
            title="选择文件",
            filetypes=[
                ("所有支持的文件", "*.pdf;*.jpg;*.jpeg;*.png"),
                ("PDF文件", "*.pdf"),
                ("图片文件", "*.jpg;*.jpeg;*.png"),
                ("所有文件", "*.*")
            ]
        )
        if files:
            self.add_files(list(files))

    def add_files(self, file_paths: List[str]):
        """添加文件"""
        for file_path in file_paths:
            file_path = file_path.strip().strip('"')
            if not os.path.exists(file_path):
                continue

            ext = os.path.splitext(file_path)[1].lower()

            if ext == '.pdf':
                if self.pdf_file:
                    if messagebox.askyesno("替换PDF", "已有PDF文件，是否替换？"):
                        self.pdf_file = file_path
                else:
                    self.pdf_file = file_path

            elif ext in ['.jpg', '.jpeg', '.png']:
                if len(self.image_files) >= 2:
                    if messagebox.askyesno("替换图片", "已有2张图片，是否替换最早的一张？"):
                        self.image_files.pop(0)
                        self.image_files.append(file_path)
                else:
                    self.image_files.append(file_path)

        self.update_display()

    def clear_files(self):
        """清除所有文件"""
        self.pdf_file = None
        self.image_files = []
        self.extracted_data = None
        self.update_display()

    def extract_data(self):
        """提取发票数据"""
        if not self.pdf_file or not PDF_AVAILABLE:
            return
            
        self.extract_btn.config(state=tk.DISABLED, text="🔄 提取中...")
        self.status_label.config(text="正在提取发票数据...")
        
        def extract_worker():
            try:
                data = self.extract_invoice_data(self.pdf_file)
                self.root.after(0, lambda: self.extract_success(data))
            except Exception as e:
                self.root.after(0, lambda: self.extract_failed(str(e)))
        
        thread = threading.Thread(target=extract_worker, daemon=True)
        thread.start()

    def extract_invoice_data(self, pdf_path: str) -> Dict[str, Any]:
        """从PDF中提取发票关键信息"""
        doc = pdfium.PdfDocument(pdf_path)
        full_text = ""
        
        for page_num in range(min(3, len(doc))):
            page = doc[page_num]
            textpage = page.get_textpage()
            text = textpage.get_text_range()
            full_text += text + "\n"
            textpage.close()
            page.close()
        
        doc.close()
        
        # 提取关键信息
        data = {
            "file_path": pdf_path,
            "file_name": os.path.basename(pdf_path),
            "extracted_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }
        
        # 发票号码
        invoice_patterns = [r'发票号码[：:\s]*(\d{8,20})', r'(\d{20})', r'(\d{12})']
        data["invoice_number"] = self.extract_by_patterns(full_text, invoice_patterns) or "未识别"
            
        # 开票日期
        date_patterns = [r'开票日期[：:\s]*(\d{4}[-年]\d{1,2}[-月]\d{1,2}日?)', r'(\d{4}[-年]\d{1,2}[-月]\d{1,2}日?)']
        raw_date = self.extract_by_patterns(full_text, date_patterns)
        if raw_date:
            data["invoice_date"] = re.sub(r'年|月', '-', raw_date).replace('日', '')
        else:
            data["invoice_date"] = "未识别"
            
        # 金额 - 提取价税合计（含税总额）
        # 使用多种策略提取，确保获取包含税额的总金额
        data["amount"] = self.extract_total_amount(full_text)
            
        # 销售方名称
        seller_patterns = [r'销售方[：:\s]*([^\n\r]+?)(?:\s|纳税人)', r'卖方[：:\s]*([^\n\r]+?)(?:\s|纳税人)']
        seller_name = self.extract_by_patterns(full_text, seller_patterns)
        data["seller_name"] = seller_name[:10] if seller_name and len(seller_name) > 3 else "未识别"
        
        return data

    def extract_by_patterns(self, text: str, patterns: List[str]) -> Optional[str]:
        """使用多个正则模式提取文本"""
        for pattern in patterns:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                return match.group(1).strip()
        return None

    def extract_total_amount(self, text: str) -> float or str:
        """
        提取发票价税合计金额（含税总额）
        使用多种策略确保提取到正确的含税总金额
        """
        # 策略1: 查找"价税合计"关键字后的金额
        # 这是最可靠的方法，因为"价税合计"是法定的含税总额
        patterns_strategy1 = [
            # 匹配: 价税合计(大写)xxx(小写)¥123.45 或 价税合计(大写)xxx(小写) ¥123.45
            r'价税合计[\s\S]{0,50}?小写[）\)]*[：:\s]*¥?\s*([\d,，]+\.?\d*)',
            # 匹配: 价税合计...¥123.45 (在价税合计后100字符内查找¥符号)
            r'价税合计[\s\S]{0,100}?¥\s*([\d,，]+\.\d{2})',
            # 匹配: 价税合计 后面直接跟数字
            r'价税合计[：:\s]+¥?\s*([\d,，]+\.\d{2})',
        ]

        for pattern in patterns_strategy1:
            match = re.search(pattern, text, re.IGNORECASE | re.DOTALL)
            if match:
                raw_amount = match.group(1)
                try:
                    clean_amount = raw_amount.replace(',', '').replace('，', '').strip()
                    amount = float(clean_amount)
                    # 金额合理性检查：一般发票金额在0.01-999999999之间
                    if 0.01 <= amount <= 999999999:
                        return amount
                except ValueError:
                    continue

        # 策略2: 查找"合计"相关字段（次优选择）
        patterns_strategy2 = [
            r'合计金额[：:\s]*¥?\s*([\d,，]+\.\d{2})',
            r'总金额[：:\s]*¥?\s*([\d,，]+\.\d{2})',
            r'应付金额[：:\s]*¥?\s*([\d,，]+\.\d{2})',
        ]

        for pattern in patterns_strategy2:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                raw_amount = match.group(1)
                try:
                    clean_amount = raw_amount.replace(',', '').replace('，', '').strip()
                    amount = float(clean_amount)
                    if 0.01 <= amount <= 999999999:
                        return amount
                except ValueError:
                    continue

        # 策略3: 查找发票中所有的金额，选择最大的一个（通常价税合计是最大值）
        # 这个策略作为最后的保底方案
        all_amounts = []
        amount_pattern = r'¥\s*([\d,，]+\.\d{2})'
        for match in re.finditer(amount_pattern, text):
            try:
                raw_amount = match.group(1)
                clean_amount = raw_amount.replace(',', '').replace('，', '').strip()
                amount = float(clean_amount)
                if 0.01 <= amount <= 999999999:
                    all_amounts.append(amount)
            except ValueError:
                continue

        if all_amounts:
            # 返回最大金额（通常就是价税合计）
            return max(all_amounts)

        return "未识别"

    def extract_success(self, data):
        """数据提取成功"""
        self.extracted_data = data
        self.extract_btn.config(state=tk.NORMAL, text="🔍 提取数据")
        self.status_label.config(text="✅ 数据提取成功！")
        self.update_display()

    def extract_failed(self, error_msg):
        """数据提取失败"""
        self.extract_btn.config(state=tk.NORMAL, text="🔍 提取数据")
        self.status_label.config(text="❌ 数据提取失败，将使用默认命名")
        self.update_display()

    def generate_smart_filename(self, data: Dict[str, Any], original_filename: str) -> str:
        """生成智能文件名"""
        base_name = os.path.splitext(original_filename)[0]
        parts = []
        
        # 添加日期
        if data.get('invoice_date') != "未识别":
            try:
                date_str = str(data['invoice_date']).replace('-', '')
                parts.append(date_str)
            except:
                pass
        
        # 添加金额
        if data.get('amount') != "未识别":
            try:
                amount_str = f"{float(data['amount']):.2f}元".replace('.00元', '元')
                parts.append(amount_str)
            except:
                pass
        
        # 添加发票号后4位
        if data.get('invoice_number') != "未识别":
            try:
                last4 = str(data['invoice_number'])[-4:]
                parts.append(f"#{last4}")
            except:
                pass
        
        if parts:
            smart_name = '_'.join(parts) + '_已合并'
        else:
            smart_name = base_name + '_已合并'
        
        return smart_name + '.pdf'

    def update_display(self):
        """更新显示"""
        # 更新文件列表
        self.file_listbox.delete(0, tk.END)
        if self.pdf_file:
            self.file_listbox.insert(tk.END, f"📄 PDF: {os.path.basename(self.pdf_file)}")
        for i, img_file in enumerate(self.image_files, 1):
            self.file_listbox.insert(tk.END, f"🖼️ 图片{i}: {os.path.basename(img_file)}")

        # 更新按钮状态
        if self.pdf_file and PDF_AVAILABLE:
            self.extract_btn.config(state=tk.NORMAL)
        else:
            self.extract_btn.config(state=tk.DISABLED)

        if self.pdf_file and len(self.image_files) == 2:
            self.merge_btn.config(state=tk.NORMAL)
            self.status_label.config(text="✅ 文件已就绪，可以合并")
        else:
            self.merge_btn.config(state=tk.DISABLED)
            if not self.pdf_file:
                self.status_label.config(text="请选择PDF文件")
            elif len(self.image_files) != 2:
                self.status_label.config(text=f"需要2张图片，当前有{len(self.image_files)}张")

        # 更新数据显示
        self.data_text.config(state=tk.NORMAL)
        self.data_text.delete(1.0, tk.END)
        
        display_text = ""
        
        if self.pdf_file:
            original_name = os.path.basename(self.pdf_file)
            
            if self.extracted_data:
                display_text += "🔍 发票数据提取结果：\n"
                display_text += "=" * 50 + "\n"
                display_text += f"📄 文件名: {self.extracted_data.get('file_name', '未知')}\n"
                display_text += f"🔢 发票号码: {self.extracted_data.get('invoice_number', '未识别')}\n"
                display_text += f"📅 开票日期: {self.extracted_data.get('invoice_date', '未识别')}\n"
                display_text += f"💰 金额: ¥{self.extracted_data.get('amount', '未识别')}\n"
                display_text += f"🏢 销售方: {self.extracted_data.get('seller_name', '未识别')}\n"
                display_text += f"⏰ 提取时间: {self.extracted_data.get('extracted_at', '')}\n\n"
                
                smart_name = self.generate_smart_filename(self.extracted_data, original_name)
                display_text += "📝 智能重命名预览：\n"
                display_text += "-" * 30 + "\n"
                display_text += f"原文件名：{original_name}\n"
                display_text += f"新文件名：{smart_name}\n"
            else:
                base_name = os.path.splitext(original_name)[0]
                default_name = base_name + "_已合并.pdf"
                display_text += "📝 默认命名预览：\n"
                display_text += f"原文件名：{original_name}\n"
                display_text += f"合并文件名：{default_name}\n\n"
                display_text += "💡 点击【提取数据】按钮启用智能命名功能"
        else:
            display_text = "💡 使用说明：\n"
            display_text += "1. 选择或拖放1个PDF发票文件\n"
            display_text += "2. 选择或拖放2张图片文件（购买记录+支付记录）\n"
            display_text += "3. 点击【提取数据】自动识别发票信息\n"
            display_text += "4. 点击【智能合并】生成最终文件\n\n"
            display_text += "✨ 功能特色：\n"
            display_text += "• 智能识别发票号码、日期、金额等关键信息\n"
            display_text += "• 根据发票内容自动生成文件名\n"
            display_text += "• 自动记录到CSV汇总文件中"
        
        self.data_text.insert(1.0, display_text)
        self.data_text.config(state=tk.DISABLED)

    def merge_files(self):
        """合并文件"""
        if not (self.pdf_file and len(self.image_files) == 2):
            messagebox.showerror("文件不完整", "需要1个PDF文件和2张图片才能合并")
            return

        # 生成文件名
        original_filename = os.path.basename(self.pdf_file)
        if self.extracted_data:
            smart_filename = self.generate_smart_filename(self.extracted_data, original_filename)
        else:
            base_name = os.path.splitext(original_filename)[0]
            smart_filename = base_name + "_已合并.pdf"

        # 弹出保存对话框
        output_path = filedialog.asksaveasfilename(
            title="保存合并后的PDF",
            defaultextension=".pdf",
            initialfile=smart_filename,
            filetypes=[("PDF文件", "*.pdf"), ("所有文件", "*.*")]
        )

        if not output_path:
            return

        self.merge_btn.config(state=tk.DISABLED, text="🔄 处理中...")
        self.status_label.config(text="正在合并文件...")

        try:
            # 创建临时文件
            temp_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
            temp_buy = tempfile.NamedTemporaryFile(suffix='.jpg', delete=False)  
            temp_pay = tempfile.NamedTemporaryFile(suffix='.jpg', delete=False)
            
            temp_pdf.close()
            temp_buy.close()
            temp_pay.close()

            # 复制文件
            shutil.copy2(self.pdf_file, temp_pdf.name)
            sorted_images = sorted(self.image_files, key=lambda x: os.path.basename(x).lower())
            shutil.copy2(sorted_images[0], temp_buy.name)
            shutil.copy2(sorted_images[1], temp_pay.name)

            # 调用合并函数
            merge_simple(temp_pdf.name, temp_buy.name, temp_pay.name, output_path)

            # 记录到CSV
            if self.extracted_data:
                self.save_to_csv(self.extracted_data, os.path.basename(output_path))

            # 清理临时文件
            for temp_file in [temp_pdf.name, temp_buy.name, temp_pay.name]:
                try:
                    os.unlink(temp_file)
                except:
                    pass

            self.merge_btn.config(state=tk.NORMAL, text="🚀 智能合并")

            # 构建成功消息
            success_msg = f"✅ 合并成功！文件已保存"
            if self.extracted_data:
                success_msg += f" | 已记录到CSV"

            self.status_label.config(text=success_msg)

            # 自动清除文件，准备处理下一张
            self.clear_files()

        except Exception as e:
            self.merge_btn.config(state=tk.NORMAL, text="🚀 智能合并")
            self.status_label.config(text="❌ 合并失败")
            messagebox.showerror("合并失败", f"合并过程中出现错误：\n{e}")

    def save_to_csv(self, data: Dict[str, Any], merged_filename: str):
        """保存到CSV文件"""
        try:
            with open(self.csv_path, 'a', encoding='utf-8-sig', newline='') as f:
                writer = csv.writer(f)
                # 发票号码添加制表符前缀，防止Excel将其转为科学计数法
                invoice_number = data.get('invoice_number', '')
                if invoice_number and invoice_number != '未识别':
                    invoice_number = f"\t{invoice_number}"

                writer.writerow([
                    invoice_number,
                    data.get('invoice_date', ''),
                    data.get('amount', ''),
                    data.get('seller_name', ''),
                    data.get('file_name', ''),
                    merged_filename,
                    datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                ])
        except Exception as e:
            print(f"保存CSV记录失败: {e}")

    def batch_extract_invoices(self):
        """批量提取发票信息"""
        if not PDF_AVAILABLE:
            messagebox.showerror("功能不可用", "批量提取功能需要pypdfium2库支持！\n\n请先安装：pip install pypdfium2")
            return

        if not EXCEL_AVAILABLE:
            messagebox.showerror("功能不可用", "批量提取功能需要openpyxl库支持！\n\n请先安装：pip install openpyxl")
            return

        # 选择文件夹
        folder_path = filedialog.askdirectory(
            title="选择包含发票PDF的文件夹"
        )

        if not folder_path:
            return

        # 禁用按钮
        self.batch_extract_btn.config(state=tk.DISABLED, text="🔄 处理中...")
        self.status_label.config(text="正在扫描文件夹...")

        # 在后台线程中处理
        def batch_worker():
            try:
                # 扫描PDF文件
                pdf_files = []
                for root, dirs, files in os.walk(folder_path):
                    for file in files:
                        if file.lower().endswith('.pdf'):
                            pdf_files.append(os.path.join(root, file))

                if not pdf_files:
                    self.root.after(0, lambda: self.batch_extract_failed("未找到PDF文件"))
                    return

                total = len(pdf_files)
                self.root.after(0, lambda: self.status_label.config(
                    text=f"找到 {total} 个PDF文件，开始提取数据...")
                )

                # 提取数据
                extracted_list = []
                success_count = 0
                fail_count = 0

                for idx, pdf_path in enumerate(pdf_files, 1):
                    try:
                        self.root.after(0, lambda i=idx, t=total, n=os.path.basename(pdf_path):
                            self.status_label.config(text=f"正在处理 {i}/{t}: {n[:30]}...")
                        )

                        data = self.extract_invoice_data(pdf_path)
                        extracted_list.append(data)
                        success_count += 1
                    except Exception as e:
                        fail_count += 1
                        print(f"提取失败 {pdf_path}: {e}")

                # 导出到Excel
                self.root.after(0, lambda: self.status_label.config(text="正在生成Excel文件..."))
                excel_path = self.export_to_excel(extracted_list, folder_path)

                # 完成
                self.root.after(0, lambda: self.batch_extract_success(
                    total, success_count, fail_count, excel_path
                ))

            except Exception as e:
                self.root.after(0, lambda: self.batch_extract_failed(str(e)))

        thread = threading.Thread(target=batch_worker, daemon=True)
        thread.start()

    def export_to_excel(self, data_list: List[Dict[str, Any]], folder_path: str) -> str:
        """导出数据到Excel文件"""
        # 创建工作簿
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "发票信息汇总"

        # 设置表头样式
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_alignment = Alignment(horizontal="center", vertical="center")

        # 表头
        headers = ['序号', '发票号码', '开票日期', '金额', '销售方名称', '文件名', '文件路径', '提取时间']
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        # 数据行样式
        data_alignment = Alignment(horizontal="left", vertical="center")
        number_alignment = Alignment(horizontal="right", vertical="center")

        # 填充数据
        for row_idx, data in enumerate(data_list, 2):
            ws.cell(row=row_idx, column=1, value=row_idx-1).alignment = number_alignment  # 序号
            ws.cell(row=row_idx, column=2, value=data.get('invoice_number', '')).alignment = data_alignment
            ws.cell(row=row_idx, column=3, value=data.get('invoice_date', '')).alignment = data_alignment

            # 金额特殊处理
            amount = data.get('amount', '')
            if isinstance(amount, (int, float)):
                ws.cell(row=row_idx, column=4, value=amount).alignment = number_alignment
            else:
                ws.cell(row=row_idx, column=4, value=str(amount)).alignment = data_alignment

            ws.cell(row=row_idx, column=5, value=data.get('seller_name', '')).alignment = data_alignment
            ws.cell(row=row_idx, column=6, value=data.get('file_name', '')).alignment = data_alignment
            ws.cell(row=row_idx, column=7, value=data.get('file_path', '')).alignment = data_alignment
            ws.cell(row=row_idx, column=8, value=data.get('extracted_at', '')).alignment = data_alignment

        # 调整列宽
        column_widths = [8, 20, 15, 12, 25, 30, 50, 20]
        for col_idx, width in enumerate(column_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

        # 保存文件
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        excel_filename = f"发票信息汇总_{timestamp}.xlsx"
        excel_path = os.path.join(folder_path, excel_filename)
        wb.save(excel_path)

        return excel_path

    def batch_extract_success(self, total: int, success: int, fail: int, excel_path: str):
        """批量提取成功"""
        self.batch_extract_btn.config(state=tk.NORMAL, text="📊 批量提取")
        self.status_label.config(text=f"✅ 批量提取完成！成功 {success}/{total} 个")

        messagebox.showinfo(
            "批量提取完成",
            f"✅ 处理完成！\n\n"
            f"总文件数：{total}\n"
            f"成功提取：{success}\n"
            f"提取失败：{fail}\n\n"
            f"Excel文件已保存：\n{os.path.basename(excel_path)}\n\n"
            f"位置：{os.path.dirname(excel_path)}"
        )

    def batch_extract_failed(self, error_msg: str):
        """批量提取失败"""
        self.batch_extract_btn.config(state=tk.NORMAL, text="📊 批量提取")
        self.status_label.config(text="❌ 批量提取失败")
        messagebox.showerror("批量提取失败", f"处理过程中出现错误：\n{error_msg}")

    def run(self):
        """运行主程序"""
        # 居中显示窗口
        self.root.update_idletasks()
        width = self.root.winfo_width()
        height = self.root.winfo_height()
        x = (self.root.winfo_screenwidth() // 2) - (width // 2)
        y = (self.root.winfo_screenheight() // 2) - (height // 2)
        self.root.geometry(f"{width}x{height}+{x}+{y}")

        # 初始化显示
        self.root.after(100, self.update_display)
        self.root.mainloop()


def main():
    """主入口"""
    try:
        if not PDF_AVAILABLE:
            messagebox.showwarning(
                "功能提示",
                "未检测到pypdfium2库，数据提取功能将不可用。\n\n"
                "安装方法：pip install pypdfium2\n\n"
                "您仍可使用基础的文件合并功能。"
            )
        
        if not DRAG_DROP_AVAILABLE:
            messagebox.showinfo(
                "功能提示", 
                "未检测到tkinterdnd2库，拖放功能不可用。\n\n"
                "请使用【选择文件】按钮选择文件。\n"
                "安装方法：pip install tkinterdnd2"
            )
            
        app = SimpleInvoiceMergerV5()
        app.run()
        return 0
        
    except Exception as e:
        messagebox.showerror("启动失败", f"程序启动失败：\n{e}")
        return 1


if __name__ == "__main__":
    sys.exit(main())